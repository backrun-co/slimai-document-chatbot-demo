import numpy as np
from docx import Document
from io import BytesIO
from openpyxl import Workbook
from pptx import Presentation

from rag import Chunk, build_context, extract_document, normalize_text, split_text


def test_normalize_text_keeps_paragraphs():
    assert normalize_text("  Xin   chào  \n\n Thế giới ") == "Xin chào\nThế giới"


def test_split_text_overlaps_and_preserves_content():
    text = " ".join(f"từ-{index}" for index in range(100))
    chunks = split_text(text, chunk_size=100, overlap=20)
    assert len(chunks) > 1
    assert all(len(chunk) <= 100 for chunk in chunks)
    assert "từ-0" in chunks[0]
    assert "từ-99" in chunks[-1]


def test_context_contains_source_metadata():
    from rag import RetrievedChunk

    result = RetrievedChunk(Chunk("Nội dung", "noi-bo.pdf", 3, 0), 0.8)
    context = build_context([result])
    assert "[Nguồn 1]" in context
    assert "noi-bo.pdf" in context
    assert "Vị trí: Trang 3" in context


def test_extract_plain_text_document():
    chunks, summary = extract_document("Nội quy làm việc từ 8 giờ.".encode(), "noi-quy.txt")
    assert chunks[0].text == "Nội quy làm việc từ 8 giờ."
    assert chunks[0].source_location == "Đoạn 1"
    assert summary == "1 đoạn văn bản"


def test_extract_csv_preserves_row_location():
    chunks, summary = extract_document(b"name,score\nAn,9\nBinh,8\n", "diem.csv")
    assert "name\tscore" in chunks[0].text
    assert chunks[0].source_location == "CSV · dòng 1–3"
    assert summary == "3 dòng CSV"


def test_extract_html_ignores_script_content():
    content = b"<h1>Chinh sach</h1><script>ignore_me()</script><p>Nghi 12 ngay.</p>"
    chunks, _ = extract_document(content, "policy.html")
    assert "Chinh sach" in chunks[0].text
    assert "ignore_me" not in chunks[0].text


def test_extract_docx_document():
    buffer = BytesIO()
    document = Document()
    document.add_paragraph("Quy trình duyệt gồm hai bước.")
    document.save(buffer)
    chunks, summary = extract_document(buffer.getvalue(), "quy-trinh.docx")
    assert "hai bước" in chunks[0].text
    assert summary == "1 đoạn Word"


def test_extract_pptx_document():
    buffer = BytesIO()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Kế hoạch Q4"
    slide.placeholders[1].text = "Tăng trưởng 20%"
    presentation.save(buffer)
    chunks, summary = extract_document(buffer.getvalue(), "ke-hoach.pptx")
    assert "Tăng trưởng 20%" in chunks[0].text
    assert chunks[0].source_location == "Slide 1"
    assert summary == "1 slide"


def test_extract_xlsx_document():
    buffer = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Doanh thu"
    sheet.append(["Tháng", "Giá trị"])
    sheet.append(["10", 120])
    workbook.save(buffer)
    chunks, summary = extract_document(buffer.getvalue(), "bao-cao.xlsx")
    assert "Tháng\tGiá trị" in chunks[0].text
    assert chunks[0].source_location == 'Sheet "Doanh thu" · dòng 1–2'
    assert summary == "1 sheet"


def test_embedding_shape_example():
    matrix = np.empty((0, 0), dtype=np.float32)
    assert matrix.size == 0

